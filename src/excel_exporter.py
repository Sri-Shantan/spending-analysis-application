from pathlib import Path
import sqlite3
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)


def _export_transactions(conn, wb):
    ws = wb.active
    ws.title = "Transactions"
    headers = ["date_of_transaction", "posted_date", "description", "merchant", "amount", "account", "transaction_type", "category", "subcategory", "source_file", "needs_review"]
    ws.append(headers)
    _style_header(ws)
    rows = conn.execute(
        "SELECT transaction_date,posted_date,description,merchant,amount,account,transaction_type,category,subcategory,source_file,needs_review FROM transactions ORDER BY transaction_date"
    ).fetchall()
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 14
    for cell in ws["E"][1:]:
        cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
    last = max(ws.max_row, 2)
    ws.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006")))
    ws.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="006100")))


def _export_statements(conn, wb):
    ws = wb.create_sheet("Statements")
    headers = [
        "statement_id", "source_file", "account", "statement_start", "statement_end",
        "statement_hash", "ingested_at", "status", "transaction_count", "needs_review_count",
        "expense_total", "income_total", "transfer_total", "adjustment_total",
    ]
    ws.append(headers)
    _style_header(ws)

    rows = conn.execute(
        """
        SELECT
            s.id, s.source_file, s.account, s.statement_start, s.statement_end,
            s.statement_hash, s.ingested_at, s.status,
            COUNT(t.id), COALESCE(SUM(t.needs_review), 0),
            COALESCE(SUM(CASE WHEN t.transaction_type = 'EXPENSE' THEN -t.amount ELSE 0 END), 0),
            COALESCE(SUM(CASE WHEN t.transaction_type = 'INCOME' THEN t.amount ELSE 0 END), 0),
            COALESCE(SUM(CASE WHEN t.transaction_type = 'TRANSFER' THEN t.amount ELSE 0 END), 0),
            COALESCE(SUM(CASE WHEN t.transaction_type = 'ADJUSTMENT' THEN t.amount ELSE 0 END), 0)
        FROM statements s
        LEFT JOIN transactions t ON t.statement_id = s.id
        GROUP BY s.id
        ORDER BY s.statement_start, s.account, s.id
        """
    ).fetchall()
    for row in rows:
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"A": 14, "B": 28, "C": 22, "D": 16, "E": 16, "F": 68, "G": 22, "H": 18, "I": 18, "J": 18, "K": 16, "L": 16, "M": 16, "N": 16}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for column in ("K", "L", "M", "N"):
        for cell in ws[column][1:]:
            cell.number_format = '$#,##0.00;[Red]-$#,##0.00'


def export(conn: sqlite3.Connection, output: str | Path) -> None:
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _export_transactions(conn, wb)
    _export_statements(conn, wb)
    wb.save(output)
